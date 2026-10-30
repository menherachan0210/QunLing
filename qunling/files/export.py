"""Artifact creation from explicit data schemas, never from executable code."""

from __future__ import annotations

import base64
import binascii
import csv
import io
import json
import os
import re
import zipfile
from pathlib import Path

from .core import (DEFAULT_LIMITS, FileError, FileLimits, dependency, digest,
                   identity, safe_member, scalar, strict_keys)
from .formats import MIMES, detect


def _text(value, limits):
    if not isinstance(value, str) or len(value) > limits.max_text_chars:
        raise FileError("invalid_schema", "文本必须为字符串且不能超过字符限制。")
    if any(ord(c) < 32 and c not in "\t\r\n" for c in value):
        raise FileError("invalid_schema", "文本包含不支持的控制字符。")
    return value


def _rows(value, limits):
    if not isinstance(value, list) or len(value) > limits.max_rows:
        raise FileError("invalid_schema", "rows 必须为受限的二维数组。")
    result, count = [], 0
    for row in value:
        if not isinstance(row, list):
            raise FileError("invalid_schema", "每行必须是数组。")
        count += len(row)
        if count > limits.max_cells:
            raise FileError("cell_limit", "输出单元格数量超过限制。")
        result.append([scalar(cell) for cell in row])
    return result


def _json(value, limits):
    # Validate structure before recursive serialization; reject custom objects and
    # cyclic graphs, non-string keys, NaN/Infinity and excessive nesting.
    stack, visited, nodes = [(value, 0)], set(), 0
    while stack:
        item, depth = stack.pop()
        nodes += 1
        if depth > 64 or nodes > limits.max_cells:
            raise FileError("invalid_schema", "数据结构过深或节点数量超过限制。")
        if isinstance(item, (dict, list)):
            if id(item) in visited:
                raise FileError("invalid_schema", "不接受循环或重复引用的数据结构。")
            visited.add(id(item))
            if isinstance(item, dict):
                if any(not isinstance(key, str) for key in item):
                    raise FileError("invalid_schema", "对象键必须为字符串。")
                stack.extend((v, depth+1) for v in item.values())
            else:
                stack.extend((v, depth+1) for v in item)
        else:
            scalar(item)
    try:
        rendered = json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False)
    except (ValueError, TypeError, RecursionError):
        raise FileError("invalid_schema", "数据不是有效 JSON。") from None
    if len(rendered) > limits.max_text_chars:
        raise FileError("text_limit", "输出数据超过字符限制。")
    return rendered


def export_file(format: str, payload: dict, output_path, *, limits: FileLimits = DEFAULT_LIMITS, staged_sources=None):
    """Create a *new* trusted-host destination; existing files are never replaced.

    Schemas:
      txt/md: {"text": str}
      csv: {"rows": [[scalar, ...], ...]}
      json: {"data": JSON value}
      docx: {"blocks": [{"text": str} | {"heading": str, "level": 1..6}
                       | {"table": [[scalar, ...], ...]}, ...]}
      xlsx: {"sheets": [{"name": str, "rows": [[scalar, ...], ...]}, ...]}
      pptx: {"slides": [{"title": str, "text": str}, ...]}
      zip: {"files": [{"name": safe relative path, "text": str}
                     | {"name": safe relative path, "base64": str}, ...]}
    XLSX strings remain strings (including leading '='); no formula execution.
    """
    if format not in {"txt", "md", "csv", "json", "docx", "xlsx", "pptx", "zip"}:
        raise FileError("unsupported_export", "尚未提供该格式的结构化导出。")
    # Bound request content independently of binary output.
    if format != "zip":
        _json(payload, limits)
    stream, warnings = io.BytesIO(), []
    try:
        if format in {"txt", "md"}:
            strict_keys(payload, {"text"}, {"text"})
            stream.write(_text(payload["text"], limits).encode("utf-8"))
        elif format == "json":
            strict_keys(payload, {"data"}, {"data"})
            stream.write((_json(payload["data"], limits) + "\n").encode("utf-8"))
        elif format == "csv":
            strict_keys(payload, {"rows"}, {"rows"})
            text = io.StringIO(newline="")
            writer = csv.writer(text)
            escaped = False
            for row in _rows(payload["rows"], limits):
                safe = []
                for cell in row:
                    if isinstance(cell, str) and cell.lstrip().startswith(("=", "+", "-", "@")):
                        cell = "'" + cell
                        escaped = True
                    safe.append(cell)
                writer.writerow(safe)
            if escaped:
                warnings.append("csv_formula_like_strings_prefixed_with_apostrophe")
            stream.write(text.getvalue().encode("utf-8-sig"))
        elif format == "docx":
            strict_keys(payload, {"blocks"}, {"blocks"})
            blocks = payload["blocks"]
            if not isinstance(blocks, list) or len(blocks) > limits.max_sections:
                raise FileError("invalid_schema", "blocks 必须为受限的正文块数组。")
            document = dependency("docx").Document()
            for block in blocks:
                strict_keys(block, {"text", "heading", "level", "table"})
                if sum(key in block for key in ("text", "heading", "table")) != 1:
                    raise FileError("invalid_schema", "正文块必须明确选择 text、heading 或 table。")
                if "text" in block:
                    strict_keys(block, {"text"})
                    document.add_paragraph(_text(block["text"], limits))
                elif "heading" in block:
                    strict_keys(block, {"heading", "level"})
                    level = block.get("level", 1)
                    if type(level) is not int or not 1 <= level <= 6:
                        raise FileError("invalid_schema", "标题等级必须为 1–6。")
                    document.add_heading(_text(block["heading"], limits), level)
                else:
                    strict_keys(block, {"table"})
                    rows = _rows(block["table"], limits)
                    width = max(map(len, rows), default=0)
                    if not rows or not width:
                        raise FileError("invalid_schema", "表格不能为空。")
                    if len(rows) * width > limits.max_cells:
                        raise FileError("cell_limit", "补齐后的表格大小超过限制。")
                    table = document.add_table(rows=len(rows), cols=width)
                    for number, row in enumerate(rows):
                        for column, value in enumerate(row):
                            table.cell(number, column).text = "" if value is None else str(value)
            document.save(stream)
        elif format == "xlsx":
            strict_keys(payload, {"sheets"}, {"sheets"})
            sheets = payload["sheets"]
            if not isinstance(sheets, list) or not 1 <= len(sheets) <= limits.max_sections:
                raise FileError("invalid_schema", "sheets 必须为非空且受限的数组。")
            library = dependency("openpyxl")
            document = library.Workbook(write_only=True)
            from openpyxl.cell import WriteOnlyCell
            names, count = set(), 0
            try:
                for sheet in sheets:
                    strict_keys(sheet, {"name", "rows"}, {"name", "rows"})
                    name = _text(sheet["name"], limits)
                    if (not name or len(name) > 31 or re.search(r"[\\/*?:\[\]]", name)
                            or name.startswith("'") or name.endswith("'") or name.casefold() in names):
                        raise FileError("invalid_schema", "工作表名称无效或重复。")
                    names.add(name.casefold())
                    rows = _rows(sheet["rows"], limits)
                    count += sum(map(len, rows))
                    if count > limits.max_cells:
                        raise FileError("cell_limit", "工作簿单元格总量超过限制。")
                    output = document.create_sheet(name)
                    for row in rows:
                        cells = []
                        for value in row:
                            if isinstance(value, str) and len(value) > 32767:
                                raise FileError("cell_limit", "XLSX 单元格文本超过 32767 字符。")
                            cell = WriteOnlyCell(output, value=value)
                            if isinstance(value, str):
                                cell.data_type = "s"
                            cells.append(cell)
                        output.append(cells)
                document.save(stream)
            finally:
                document.close()
        elif format == "pptx":
            strict_keys(payload, {"slides"}, {"slides"})
            slides = payload["slides"]
            if not isinstance(slides, list) or not 1 <= len(slides) <= limits.max_sections:
                raise FileError("invalid_schema", "slides 必须为非空且受限的数组。")
            document = dependency("pptx").Presentation()
            for item in slides:
                strict_keys(item, {"title", "text"}, {"title", "text"})
                slide = document.slides.add_slide(document.slide_layouts[1])
                slide.shapes.title.text = _text(item["title"], limits)
                slide.placeholders[1].text = _text(item["text"], limits)
            document.save(stream)
        else:
            strict_keys(payload, {"files"}, {"files"})
            files = payload["files"]
            if not isinstance(files, list) or len(files) > limits.max_entries:
                raise FileError("invalid_schema", "files 必须为受限的成员数组。")
            names, total = set(), 0
            # STORED keeps artifact validation deterministic even for highly
            # repetitive content, and avoids exporting compression bombs.
            with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_STORED) as archive:
                for entry in files:
                    strict_keys(entry, {"name", "text", "base64", "staged"} if staged_sources else {"name", "text", "base64"}, {"name"})
                    if sum(key in entry for key in ("text", "base64", "staged")) != 1:
                        raise FileError("invalid_schema", "成员必须选择 text 或 base64 内容之一。")
                    name = safe_member(entry["name"])
                    if name.endswith("/") or identity(name) in names:
                        raise FileError("ambiguous_member", "输出成员名称重复或不是普通文件。")
                    names.add(identity(name))
                    if "staged" in entry:
                        source = (staged_sources or {}).get(entry["staged"])
                        if source is None or source.is_symlink() or not source.is_file():
                            raise FileError("invalid_schema", "打包素材不可用。")
                        size = source.stat().st_size
                        total += size
                        if size > limits.max_member_bytes or total > limits.max_output_bytes:
                            raise FileError("output_limit", "输出压缩包超过大小限制。")
                        # zipfile streams the verified private file in chunks.
                        archive.write(source, arcname=name)
                        continue
                    elif "text" in entry:
                        raw = _text(entry["text"], limits).encode("utf-8")
                    else:
                        encoded = entry["base64"]
                        if not isinstance(encoded, str) or len(encoded) > ((limits.max_member_bytes + 2)//3)*4:
                            raise FileError("member_limit", "成员内容超过限制。")
                        try:
                            raw = base64.b64decode(encoded, validate=True)
                        except (ValueError, binascii.Error):
                            raise FileError("invalid_schema", "成员 base64 内容无效。") from None
                    total += len(raw)
                    if len(raw) > limits.max_member_bytes or total > limits.max_output_bytes:
                        raise FileError("output_limit", "输出压缩包超过大小限制。")
                    archive.writestr(name, raw)
        raw = stream.getvalue()
        if len(raw) > limits.max_output_bytes:
            raise FileError("output_limit", "生成文件超过输出大小限制。")
        detected, _, _ = detect(raw, "artifact." + format, limits)
        text_override = format in {"txt", "md"} and detected in {"txt", "md", "json", "source"}
        if detected != format and not text_override:
            # Empty text is TXT; empty Markdown and JSON scalar detection still
            # follow their explicit filename subtype.
            raise FileError("invalid_output", "生成文件未通过格式检查。")
    except FileError:
        raise
    except Exception:
        raise FileError("export_failed", "生成文件失败，输入结构或内容不受支持。") from None
    destination = Path(output_path)
    try:
        # Parent directory is chosen by the host; no model-provided destination.
        with destination.open("xb") as output:
            try:
                output.write(raw)
                output.flush()
                os.fsync(output.fileno())
            except BaseException:
                output.close()
                destination.unlink(missing_ok=True)
                raise
    except FileExistsError:
        raise FileError("output_exists", "目标文件已存在，未覆盖。") from None
    except OSError:
        raise FileError("output_unavailable", "目标文件无法写入。") from None
    return {"format": format, "mime": MIMES[format], "size": len(raw),
            "sha256": digest(raw), "warnings": warnings}
