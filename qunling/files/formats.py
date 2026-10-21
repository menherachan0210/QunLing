"""Real-content detection and parsers producing navigable document sections."""

from __future__ import annotations

import csv
import io
import json
import math
import re
import xml.parsers.expat
from datetime import date, datetime, time
from pathlib import Path

from .archives import ZipArchive, archive_type, open_archive
from .core import FileError, FileLimits, dependency

MIMES = {
    "txt": "text/plain", "md": "text/markdown", "csv": "text/csv",
    "tsv": "text/tab-separated-values", "json": "application/json",
    "source": "text/plain", "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "zip": "application/zip", "rar": "application/vnd.rar", "7z": "application/x-7z-compressed",
    "png": "image/png", "jpeg": "image/jpeg", "gif": "image/gif", "webp": "image/webp",
    "unknown": "application/octet-stream",
}
SOURCE_SUFFIXES = {
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".c", ".cpp", ".h", ".hpp",
    ".cs", ".go", ".rs", ".rb", ".php", ".sh", ".ps1", ".bat", ".cmd", ".sql",
    ".yaml", ".yml", ".toml", ".ini", ".xml", ".html", ".css", ".vue", ".svelte",
    ".log", ".conf", ".cfg", ".r", ".swift", ".kt", ".lua", ".tex",
}


def decode_text(raw: bytes) -> tuple[str, str]:
    if raw.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        raise FileError("unsupported_encoding", "暂不支持 UTF-32 文本。")
    encodings = (["utf-16"] if raw.startswith((b"\xff\xfe", b"\xfe\xff"))
                 else ["utf-8-sig", "gb18030"])
    for encoding in encodings:
        try:
            text = raw.decode(encoding, errors="strict")
        except UnicodeError:
            continue
        if any(ord(char) < 32 and char not in "\t\r\n" for char in text):
            raise FileError("binary_file", "文件包含二进制或控制数据，不能作为文本读取。")
        return text, encoding
    raise FileError("unsupported_encoding", "无法可靠识别文本编码。")


def xml_guard(raw: bytes, limits: FileLimits):
    if len(raw) > limits.max_xml_bytes:
        raise FileError("xml_limit", "文档 XML 成员超过限制。")
    parser = xml.parsers.expat.ParserCreate()
    depth, nodes = 0, 0
    def start(*_):
        nonlocal depth, nodes
        depth += 1
        nodes += 1
        if depth > limits.max_xml_depth or nodes > limits.max_xml_nodes:
            raise FileError("xml_limit", "文档 XML 结构超过限制。")
    def end(*_):
        nonlocal depth
        depth -= 1
    def doctype(*_):
        raise FileError("unsafe_xml", "不解析含 DTD 或实体声明的文档。")
    parser.StartElementHandler = start
    parser.EndElementHandler = end
    parser.StartDoctypeDeclHandler = doctype
    parser.EntityDeclHandler = doctype
    parser.ExternalEntityRefHandler = doctype
    try:
        parser.Parse(raw, True)
    except FileError:
        raise
    except xml.parsers.expat.ExpatError:
        raise FileError("invalid_document", "文档 XML 损坏。") from None


def office_preflight(archive: ZipArchive, limits: FileLimits):
    total, warnings = 0, set()
    for entry in archive.entries:
        name = entry["path"]
        lower = name.lower()
        if lower.endswith((".xml", ".rels")):
            total += entry["size"]
            if total > limits.max_xml_total_bytes:
                raise FileError("xml_limit", "文档 XML 总大小超过限制。")
            body = archive.read(name)
            xml_guard(body, limits)
            if lower.endswith(".rels") and re.search(
                rb'TargetMode\s*=\s*["\']External["\']', body, re.I
            ):
                warnings.add("external_links_not_followed")
        if "vbaproject" in lower or "/embeddings/" in lower or lower.endswith((".exe", ".dll")):
            warnings.add("embedded_content_not_executed")
    return sorted(warnings)


def detect(raw: bytes, name: str, limits: FileLimits, *, password: str | None = None):
    kind = archive_type(raw)
    if kind == "zip":
        archive = ZipArchive(raw, limits, password=password)
        if archive.encrypted:
            return "zip", [], archive
        names = {item["path"] for item in archive.entries}
        office = [fmt for fmt, marker in (
            ("docx", "word/document.xml"), ("xlsx", "xl/workbook.xml"),
            ("pptx", "ppt/presentation.xml")) if marker in names]
        if office and "[Content_Types].xml" in names:
            if len(office) != 1:
                raise FileError("ambiguous_format", "文档混合了多个 Office 格式。")
            warnings = office_preflight(archive, limits)
            return office[0], warnings, archive
        return "zip", [], archive
    if kind:
        return ("rar" if kind.startswith("Rar") else kind), [], None
    for prefix, fmt in ((b"%PDF-", "pdf"), (b"\x89PNG\r\n\x1a\n", "png"),
                        (b"\xff\xd8\xff", "jpeg"), (b"GIF87a", "gif"), (b"GIF89a", "gif")):
        if raw.startswith(prefix):
            return fmt, [], None
    if raw.startswith(b"RIFF") and raw[8:12] == b"WEBP":
        return "webp", [], None
    # Executables and OLE (legacy Office/encrypted OOXML) must never fall through
    # to permissive legacy-text decoding.
    if raw.startswith((b"MZ", b"\x7fELF", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")):
        return "unknown", ["binary_or_legacy_format"], None
    try:
        text, _ = decode_text(raw)
    except FileError:
        return "unknown", [], None
    suffix = Path(name).suffix.lower()
    if suffix == ".json" or text.lstrip().startswith(("{", "[")):
        try:
            json.loads(text, parse_constant=lambda _: (_ for _ in ()).throw(ValueError()))
            return "json", [], None
        except (ValueError, RecursionError):
            if suffix == ".json":
                return "txt", ["invalid_json"], None
    if suffix in {".csv", ".tsv"}:
        return suffix[1:], ["text_subtype_from_filename"], None
    if suffix in {".md", ".markdown"}:
        return "md", ["text_subtype_from_filename"], None
    return ("source" if suffix in SOURCE_SUFFIXES else "txt"), [], None


def _checked(sections, limits):
    if len(sections) > limits.max_sections:
        raise FileError("section_limit", "文档页数或结构数量超过限制。")
    count = sum(len(json.dumps(s, ensure_ascii=False, default=str)) for s in sections)
    if count > limits.max_text_chars:
        raise FileError("text_limit", "文档提取内容超过限制，请拆分文件。")
    return sections


def _table_rows(rows, limits):
    result, cells = [], 0
    for index, row in enumerate(rows, 1):
        if index > limits.max_rows:
            raise FileError("row_limit", "表格行数超过限制。")
        values = list(row)
        cells += len(values)
        if cells > limits.max_cells:
            raise FileError("cell_limit", "表格单元格数量超过限制。")
        result.append({"row": index, "cells": values})
    return result


def _cell_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float) and math.isfinite(value):
        return value
    raise FileError("unsupported_cell", "工作表中存在无法转换为标准数据的单元格。")


def parse(raw, fmt, limits: FileLimits, archive=None, *, password: str | None = None):
    """Return sections and extra warnings. No links/macros/formulas execute."""
    warnings = []
    if fmt in {"txt", "md", "source", "json"}:
        text, encoding = decode_text(raw)
        if len(text) > limits.max_text_chars:
            raise FileError("text_limit", "文本超过提取字符限制。")
        sections = [{"id": f"text:{n // 4000 + 1}", "title": f"字符 {n + 1}–{min(n+4000,len(text))}",
                     "text": text[n:n + 4000], "start_char": n, "encoding": encoding}
                    for n in range(0, len(text), 4000)]
        return _checked(sections or [{"id": "text:1", "title": "空文件", "text": ""}], limits), warnings
    if fmt in {"csv", "tsv"}:
        text, encoding = decode_text(raw)
        if len(text) > limits.max_text_chars:
            raise FileError("text_limit", "表格文本超过提取字符限制。")
        try:
            rows = _table_rows(csv.reader(io.StringIO(text), delimiter="\t" if fmt == "tsv" else ",",
                                          strict=True), limits)
        except csv.Error:
            raise FileError("invalid_document", "CSV/TSV 数据损坏或字段过长。") from None
        return _checked([{"id": "table:1", "title": "数据表", "rows": rows,
                          "encoding": encoding}], limits), warnings
    if fmt in {"zip", "rar", "7z"}:
        opened = archive or open_archive(raw, limits, password=password)
        return [{"id": "archive:1", "title": "压缩包目录", "entries": opened.entries}], warnings
    if fmt in {"png", "jpeg", "gif", "webp"}:
        return [{"id": "image:1", "title": "图片", "needs_vision": True}], warnings
    if fmt == "pdf":
        library = dependency("pypdf")
        try:
            reader = library.PdfReader(io.BytesIO(raw), strict=True)
            if reader.is_encrypted:
                raise FileError("encrypted_file", "暂不读取加密 PDF。")
            if len(reader.pages) > limits.max_sections:
                raise FileError("section_limit", "PDF 页数超过限制。")
            sections, chars = [], 0
            for number, page in enumerate(reader.pages, 1):
                # Validate finite, reasonable page geometry before any future render.
                box = page.mediabox
                if not 0 < float(box.width) <= 20_000 or not 0 < float(box.height) <= 20_000:
                    raise FileError("page_limit", "PDF 页面尺寸超出限制。")
                text = page.extract_text() or ""
                chars += len(text)
                if chars > limits.max_text_chars:
                    raise FileError("text_limit", "PDF 提取文本超过限制。")
                sections.append({"id": f"page:{number}", "page": number,
                                 "title": f"第 {number} 页", "text": text,
                                 "needs_ocr": not bool(text.strip())})
            return _checked(sections, limits), ["pdf_active_content_not_executed"]
        except FileError:
            raise
        except Exception:
            raise FileError("invalid_document", "PDF 无法可靠解析。") from None
    if fmt == "docx":
        docx = dependency("docx")
        try:
            document = docx.Document(io.BytesIO(raw))
            # Preserve body paragraph/table order; headers/footnotes are explicit
            # omissions, not silently conflated with complete document text.
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            sections = []
            for child in document.element.body:
                tag = child.tag.rsplit("}", 1)[-1]
                if tag not in {"p", "tbl"}:
                    continue
                number = len(sections) + 1
                section = {"id": f"block:{number}", "title": f"正文块 {number}"}
                if tag == "p":
                    section["text"] = Paragraph(child, document).text
                else:
                    table = Table(child, document)
                    section["rows"] = _table_rows(
                        ([cell.text for cell in row.cells] for row in table.rows), limits
                    )
                sections.append(section)
                _checked(sections, limits)
            return sections, ["body_only_headers_footnotes_not_extracted"]
        except FileError:
            raise
        except Exception:
            raise FileError("invalid_document", "DOCX 无法可靠解析。") from None
    if fmt == "xlsx":
        library = dependency("openpyxl")
        formulas = cached = None
        try:
            formulas = library.load_workbook(io.BytesIO(raw), read_only=True, data_only=False,
                                             keep_links=False, keep_vba=False)
            cached = library.load_workbook(io.BytesIO(raw), read_only=True, data_only=True,
                                           keep_links=False, keep_vba=False)
            if len(formulas.worksheets) > limits.max_sections:
                raise FileError("section_limit", "工作表数量超过限制。")
            sections, total = [], 0
            for sheet, values in zip(formulas.worksheets, cached.worksheets):
                if (sheet.max_row or 0) > limits.max_rows:
                    raise FileError("row_limit", "工作表声明行数超过限制。")
                if (sheet.max_row or 0) * (sheet.max_column or 0) > limits.max_cells:
                    raise FileError("cell_limit", "工作表声明单元格数量超过限制。")
                rows = []
                for index, (row, value_row) in enumerate(zip(sheet.iter_rows(), values.iter_rows()), 1):
                    if index > limits.max_rows:
                        raise FileError("row_limit", "工作表实际行数超过限制。")
                    cells = []
                    for cell, cached_cell in zip(row, value_row):
                        total += 1
                        if total > limits.max_cells:
                            raise FileError("cell_limit", "工作簿单元格总量超过限制。")
                        value = cell.value
                        if cell.data_type == "f":
                            cached_value = _cell_value(cached_cell.value)
                            formula = value if isinstance(value, str) else getattr(value, "text", None)
                            if formula is not None and not isinstance(formula, str):
                                raise FileError("unsupported_cell", "公式结构无法可靠读取。")
                            item = {"formula": formula, "cached_value": cached_value,
                                    "cache_available": cached_value is not None}
                            if not isinstance(value, str):
                                item["formula_type"] = type(value).__name__
                        else:
                            item = {"value": _cell_value(value)}
                        cells.append(item)
                    rows.append({"row": index, "cells": cells})
                sections.append({"id": f"sheet:{len(sections)+1}", "title": sheet.title,
                                 "sheet": sheet.title, "rows": rows})
                _checked(sections, limits)
            return sections, ["formulas_not_evaluated_cached_values_may_be_stale"]
        except FileError:
            raise
        except Exception:
            raise FileError("invalid_document", "XLSX 无法可靠解析。") from None
        finally:
            if formulas is not None:
                formulas.close()
            if cached is not None:
                cached.close()
    if fmt == "pptx":
        library = dependency("pptx")
        try:
            document = library.Presentation(io.BytesIO(raw))
            if len(document.slides) > limits.max_sections:
                raise FileError("section_limit", "幻灯片数量超过限制。")
            sections = []
            def texts(shapes):
                for shape in shapes:
                    if getattr(shape, "has_text_frame", False):
                        yield shape.text
                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            yield "\t".join(cell.text for cell in row.cells)
                    if hasattr(shape, "shapes"):
                        yield from texts(shape.shapes)
            for number, slide in enumerate(document.slides, 1):
                sections.append({"id": f"slide:{number}", "page": number,
                                 "title": f"幻灯片 {number}", "text": "\n".join(texts(slide.shapes))})
                _checked(sections, limits)
            return sections, ["slide_text_only_images_charts_notes_not_extracted"]
        except FileError:
            raise
        except Exception:
            raise FileError("invalid_document", "PPTX 无法可靠解析。") from None
    raise FileError("unsupported_format", "暂不支持解析这种文件；可以保留和转发原文件。")
