import asyncio
from dataclasses import replace
import io

import pytest

from qunling.assets import AssetStore
from qunling.cache import JsonCache
from qunling.files import inspect_file, read_file, search_file
from qunling.files.cache import exchange
from qunling.files.core import DEFAULT_LIMITS
from qunling.processing import run_operation
from qunling.scope import Scope


def document():
    from docx import Document
    doc = Document()
    doc.add_paragraph("First section with reusable evidence")
    doc.add_paragraph("Second section")
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_scoped_cache_expiry_and_total_budget(tmp_path):
    now = [100]
    cache = JsonCache(tmp_path / "cache.db", clock=lambda: now[0], ttl=10,
                      max_items=2, max_bytes=100, max_item_bytes=60)
    assert cache.put("one", "a", {"text": "hello"})
    assert cache.get("two", "a") is None
    assert not cache.put("one", "big", "x" * 100)
    assert cache.put("one", "b", 2)
    assert cache.put("one", "c", 3)
    assert cache.get("one", "a") is None
    now[0] += 11
    assert cache.get("one", "c") is None


def test_inspect_read_search_share_parse_but_not_limits_or_passwords(tmp_path, monkeypatch):
    source = tmp_path / "sample.docx"
    source.write_bytes(document())
    with exchange() as first:
        inspect_file(source)
    assert first["result"] and not first["hit"]
    import qunling.files.formats as formats
    original = formats.parse
    calls = []
    def counted(*args, **kwargs):
        calls.append(True)
        return original(*args, **kwargs)
    monkeypatch.setattr(formats, "parse", counted)
    with exchange(first["result"]) as second:
        assert read_file(source)["sections"]
        assert search_file(source, "evidence")["hits"]
    assert second["hit"] and not calls
    with exchange(first["result"]) as third:
        inspect_file(source, limits=replace(DEFAULT_LIMITS, max_sections=99))
    assert not third["hit"] and len(calls) == 1
    with exchange(first["result"]) as fourth:
        inspect_file(source, password="private")
    assert not fourth["hit"] and fourth["result"] is None and len(calls) == 2


@pytest.mark.asyncio
async def test_real_worker_reuses_structure_after_authorized_asset_read(tmp_path):
    store = AssetStore(tmp_path / "assets")
    scope = Scope("workspace", "bot", "group:1", "member")
    asset = store.put(scope, "sample.docx", document(), origin="generated")
    first = await run_operation(store, scope, "inspect", {"asset_ref": asset["asset_ref"]})
    second = await run_operation(store, scope, "search", {"asset_ref": asset["asset_ref"], "query": "evidence"})
    assert not first["parse_cache_hit"]
    assert second["parse_cache_hit"] and second["hits"]
